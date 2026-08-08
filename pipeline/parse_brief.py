#!/usr/bin/env python3
"""기획안 xlsx → brief.json

양식 (디자인23 기준, docs/rules.md §6-1 참조):
  - Main sheet '기획안': B=섹션그룹 C=# D=섹션명 E=구분 F=카피초안 G=기획
    H=추가코멘트 I=디자인가이드 J=수정1 K=수정2 L=방어/거절멘트
    M~P = 우측 데이터 테이블([S0] 제품 스펙 같은 앵커 라벨로 시작)
  - 부가 sheet: 리서치 / 경쟁사 / 컨셉 / 줄글에세이 / FAQ

이 스크립트는 **기계적 추출만** 한다. 최종 카피 확정(수정 수용/방어 반영),
장면 분해, 레이아웃 선택은 에이전트(Claude)의 몫 — 결과는 raw 그대로 담는다.

사용: python3 pipeline/parse_brief.py <기획안.xlsx> [out.json]
"""
import json
import re
import sys

import openpyxl


def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return str(v).strip() if v is not None else None


def parse_main(ws):
    """'기획안' sheet → scene 목록 + 사이드 테이블 + 하단 배너."""
    scenes = []
    cur = None
    group = None
    banner_start = None

    for r in range(4, ws.max_row + 1):
        b = cell(ws, r, 2)
        if b and b.startswith("■"):  # '■ 하단 배너 기획'
            banner_start = r
            break
        num, name, kind = cell(ws, r, 3), cell(ws, r, 4), cell(ws, r, 5)
        if b:
            group = b
        if num:  # 새 장면 시작
            if cur:
                scenes.append(cur)
            cur = {
                "no": int(num),
                "group": group,
                "name": name,
                "rows": [],
            }
        if cur is None:
            continue
        row = {}
        for col, key in [(5, "kind"), (6, "copy"), (7, "plan"), (8, "comment"),
                         (9, "design_guide"), (10, "rev1"), (11, "rev2"), (12, "defense")]:
            v = cell(ws, r, col)
            if v:
                row[key] = v
        if row:
            cur["rows"].append(row)
    if cur:
        scenes.append(cur)

    # 우측 사이드 테이블: M열(13)에서 [Sx]/[Mx] 앵커를 찾아 아래로 긁는다
    tables = []
    for r in range(1, ws.max_row + 1):
        m = cell(ws, r, 13)
        if m and re.match(r"^\[[SM]\d", m):
            rows = []
            for rr in range(r + 1, ws.max_row + 1):
                vals = [cell(ws, rr, c) for c in range(13, 17)]
                if not any(vals):
                    break
                if vals[0] and re.match(r"^\[[SM]\d", vals[0]):
                    break
                rows.append([v for v in vals if v is not None])
            # 앵커 행 자체에 값이 더 있으면 (예: [S0] 옆 헤더) 포함
            head = [cell(ws, r, c) for c in range(14, 17)]
            tables.append({"id": m, "anchor_row": r,
                           "header": [v for v in head if v],
                           "rows": rows})
    banners = []
    if banner_start:
        for r in range(banner_start + 2, ws.max_row + 1):
            b = cell(ws, r, 2)
            if not b or b in ("구분",):
                continue
            banners.append({
                "type": b,
                "copy": cell(ws, r, 3),
                "condition": cell(ws, r, 7),
                "design_guide": cell(ws, r, 9),
            })
    return scenes, tables, banners


def parse_grid(ws):
    """부가 sheet를 행 단위 리스트로 (빈 셀 제거)."""
    out = []
    for row in ws.iter_rows():
        vals = [str(c.value).strip() for c in row if c.value is not None]
        if vals:
            out.append(vals)
    return out


def main():
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "brief.json"
    wb = openpyxl.load_workbook(src, data_only=True)

    main_ws = wb["기획안"]
    scenes, tables, banners = parse_main(main_ws)

    brief = {
        "source": src.split("/")[-1],
        "title": cell(main_ws, 1, 2),
        "scenes": scenes,
        "side_tables": tables,
        "banners": banners,
        "aux": {name: parse_grid(wb[name])
                for name in wb.sheetnames if name != "기획안"},
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(brief, f, ensure_ascii=False, indent=1)
    print(f"{out}: scenes={len(scenes)} tables={len(tables)} banners={len(banners)} "
          f"aux={list(brief['aux'])}")


if __name__ == "__main__":
    main()
